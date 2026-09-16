"""
Document parser — extract plain text from PDF and DOCX resumes.
"""

from __future__ import annotations

import io
import re
import unicodedata
from pathlib import Path
from typing import BinaryIO, Union

from pypdf import PdfReader
from pypdf.errors import PdfReadError
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def _normalize_whitespace(text: str) -> str:
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u00ad", "")
    text = re.sub(r"[\u200b\u200c\u200d\u200e\u200f\ufeff\u202a-\u202e]", "", text)
    text = re.sub(r"[\r\t]+", " ", text)
    text = re.sub(r" +", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(source: Union[str, Path, BinaryIO]) -> str:
    try:
        if isinstance(source, (str, Path)):
            reader = PdfReader(str(source))
        else:
            reader = PdfReader(io.BytesIO(source.read()))
    except PdfReadError as e:
        raise ValueError(f"Could not read PDF: {e}") from e
    except Exception as e:
        raise ValueError(f"Unexpected error reading PDF: {e}") from e

    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")

    text = _normalize_whitespace("\n".join(pages))
    if not text:
        raise ValueError("PDF contains no extractable text (may be a scanned image).")
    return text


def parse_docx(source: Union[str, Path, BinaryIO]) -> str:
    try:
        if isinstance(source, (str, Path)):
            doc = Document(str(source))
        else:
            doc = Document(io.BytesIO(source.read()))
    except PackageNotFoundError as e:
        raise ValueError(f"Not a valid DOCX file: {e}") from e
    except Exception as e:
        raise ValueError(f"Unexpected error reading DOCX: {e}") from e

    chunks = []
    for para in doc.paragraphs:
        if para.text:
            chunks.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    chunks.append(cell_text)

    text = _normalize_whitespace("\n".join(chunks))
    if not text:
        raise ValueError("DOCX contains no extractable text.")
    return text


def parse_document(source: Union[str, Path, BinaryIO], filename: str | None = None) -> str:
    if filename is None and isinstance(source, (str, Path)):
        filename = str(source)
    if filename is None:
        raise ValueError("Cannot determine file type: provide `filename` or a file path.")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    return parse_pdf(source) if suffix == ".pdf" else parse_docx(source)